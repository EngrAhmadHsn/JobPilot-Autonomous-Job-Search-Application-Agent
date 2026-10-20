"""WS-05 tests: the deterministic ATS analyzer (src/resume/ats.py).

The analyzer has no model call, so every assertion here is arithmetic: what
the report says must follow from the inputs. The "deliberately bad" DOCX
fixture exercises the format-risk detection with a document that has a layout
table, a text box, an image, missing headings, an empty section, three font
sizes and no contact line.
"""

from __future__ import annotations

import base64
from pathlib import Path

import pytest

from src.jobs.contracts import JobListing
from src.resume import parser
from src.resume.ats import (
    ATS_WEIGHTS,
    FORMAT_RISK_PENALTY,
    NEUTRAL_FORMAT_SAFETY,
    NEUTRAL_KEYWORD_SCORE,
    NEUTRAL_TITLE_ALIGNMENT,
    REQUIRED_SECTIONS,
    STANDARD_HEADINGS,
    analyze_ats,
)
from src.resume.profile import (
    CandidateProfile,
    EducationEntry,
    ExperienceEntry,
    PersonalInformation,
)
from src.resume.tailor import TailoredExperience, TailoredResume

PROFILE = CandidateProfile(
    personal_information=PersonalInformation(
        name="Ahmad Hassan",
        title="AI Engineer",
        email="ahmad@example.com",
        phone="+49 30 1234567",
        location="Berlin, Germany",
    ),
    summary="AI engineer building agent systems.",
    skills=["python", "langgraph", "fastapi", "docker"],
    education=[EducationEntry(degree="BSc Computer Science", institution="TU Berlin")],
    experience=[ExperienceEntry(title="AI Engineer", company="Agentic Labs", years=5)],
)

TAILORED = TailoredResume(
    summary="AI engineer with 5 years building agent pipelines in Python.",
    highlighted_skills=["python", "langgraph", "fastapi"],
    experience=[
        TailoredExperience(
            company="Agentic Labs",
            title="AI Engineer",
            bullets=["Built agent pipelines with LangGraph", "Shipped REST APIs with FastAPI"],
        )
    ],
    projects=[],
    cover_letter="Dear hiring team, ...",
)

JOB = JobListing(
    id="greenhouse:abc",
    source="greenhouse",
    title="AI Engineer",
    company="Acme AI",
    required_skills=["python", "langgraph"],
    preferred_skills=["fastapi"],
)

# A 1x1 transparent PNG, enough for python-docx to embed a real image.
_ONE_PIXEL_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)


def _bad_docx(path: Path) -> Path:
    """A resume with every format risk the analyzer knows about."""
    from docx import Document
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    from docx.shared import Pt

    document = Document()
    document.add_paragraph("Jane Doe")  # no email/phone -> no contact line
    document.add_paragraph("Engineer")
    document.add_paragraph("Professional Summary")
    document.add_paragraph("Summary text")
    document.add_paragraph("Technical Skills")
    document.add_paragraph("Python, LangGraph")

    # Three distinct body font sizes.
    for size in (10, 12, 14):
        paragraph = document.add_paragraph(f"Body text at {size}pt")
        for run in paragraph.runs:
            run.font.size = Pt(size)

    # A layout table.
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Left"
    table.cell(0, 1).text = "Right"

    # A text box (raw XML: python-docx has no API for it).
    document.element.body.append(
        parse_xml(
            f"<w:p {nsdecls('w')}><w:r><w:pict><w:txbxContent><w:p><w:r>"
            "<w:t>hidden text</w:t></w:r></w:p></w:txbxContent></w:pict></w:r></w:p>"
        )
    )

    # An image.
    image_path = path.with_suffix(".png")
    image_path.write_bytes(_ONE_PIXEL_PNG)
    document.add_picture(str(image_path), width=1000000)

    # An empty section: a heading followed by another heading.
    document.add_paragraph("Certifications")
    document.add_paragraph("Achievements")

    document.save(str(path))
    return path


def _clean_docx(path: Path) -> Path:
    """A minimal, ATS-friendly resume: headings, one font, contact line."""
    from docx import Document

    document = Document()
    document.add_paragraph("Ahmad Hassan")
    document.add_paragraph("AI Engineer")
    document.add_paragraph("ahmad@example.com | +49 30 1234567 | Berlin, Germany")
    document.add_paragraph("Professional Summary")
    document.add_paragraph("AI engineer with 5 years building agent pipelines in Python.")
    document.add_paragraph("Technical Skills")
    document.add_paragraph("Languages: Python, LangGraph, FastAPI")
    document.add_paragraph("Experience")
    document.add_paragraph("AI Engineer | Agentic Labs")
    document.add_paragraph("Built agent pipelines with LangGraph", style="List Bullet")
    document.add_paragraph("Education")
    document.add_paragraph("BSc Computer Science")
    document.save(str(path))
    return path


def test_standard_headings_stay_in_sync_with_the_parser() -> None:
    assert STANDARD_HEADINGS == parser._SECTION_HEADINGS


def test_keyword_coverage_counts_hits_and_misses() -> None:
    report = analyze_ats(TAILORED, PROFILE, JOB)
    coverage = report.keyword_coverage
    assert coverage.required_hit == ["python", "langgraph"]
    assert coverage.required_missing == []
    assert coverage.preferred_hit == ["fastapi"]
    assert coverage.preferred_missing == []
    assert coverage.required_coverage == 1.0
    assert coverage.preferred_coverage == 1.0


def test_keyword_coverage_partial_and_none() -> None:
    job = JobListing(
        id="j",
        source="greenhouse",
        title="AI Engineer",
        company="Acme",
        required_skills=["python", "kubernetes", "terraform"],
        preferred_skills=["fastapi"],
    )
    report = analyze_ats(TAILORED, PROFILE, job)
    coverage = report.keyword_coverage
    assert coverage.required_hit == ["python"]
    assert coverage.required_missing == ["kubernetes", "terraform"]
    assert coverage.preferred_hit == ["fastapi"]
    assert coverage.required_coverage == pytest.approx(1 / 3)
    assert coverage.preferred_coverage == 1.0


def test_preferred_keywords_count_half_a_required_one() -> None:
    """0/2 required and 1/1 preferred -> 0.5 / 1.5 of the keyword sub-score."""
    job = JobListing(
        id="j",
        source="greenhouse",
        title="AI Engineer",
        company="Acme",
        required_skills=["kubernetes", "terraform"],
        preferred_skills=["fastapi"],
    )
    report = analyze_ats(TAILORED, PROFILE, job)
    expected = 100.0 * (0.5 * 1.0) / (1.0 + 0.5)
    assert report.score == pytest.approx(
        ATS_WEIGHTS["keywords"] * expected
        + ATS_WEIGHTS["sections"] * 100.0
        + ATS_WEIGHTS["title_alignment"] * 100.0
        + ATS_WEIGHTS["format_safety"] * NEUTRAL_FORMAT_SAFETY,
        abs=0.05,
    )


def test_score_is_the_weighted_sum_of_its_components() -> None:
    report = analyze_ats(TAILORED, PROFILE, JOB)
    assert ATS_WEIGHTS["keywords"] + ATS_WEIGHTS["sections"] + ATS_WEIGHTS[
        "title_alignment"
    ] + ATS_WEIGHTS["format_safety"] == pytest.approx(1.0)
    assert 0.0 <= report.score <= 100.0
    assert report.score == pytest.approx(
        ATS_WEIGHTS["keywords"] * 100.0
        + ATS_WEIGHTS["sections"] * 100.0
        + ATS_WEIGHTS["title_alignment"] * 100.0
        + ATS_WEIGHTS["format_safety"] * NEUTRAL_FORMAT_SAFETY,
        abs=0.05,
    )


def test_neutral_values_are_used_when_evidence_is_missing() -> None:
    """A job with no skills and a candidate with no title are unknown, not bad."""
    bare_job = JobListing(id="j", source="greenhouse", title="", company="")
    bare_profile = CandidateProfile(
        personal_information=PersonalInformation(name="Ahmad Hassan"),
        skills=[],
    )
    report = analyze_ats(
        TailoredResume(
            summary="",
            highlighted_skills=[],
            experience=[],
            projects=[],
            cover_letter="",
        ),
        bare_profile,
        bare_job,
    )
    assert report.title_alignment == NEUTRAL_TITLE_ALIGNMENT
    assert report.keyword_coverage.required_coverage == 0.0
    # The keyword sub-score is unknown, not zero and not perfect; no section
    # is present either, and nothing was observed about the format.
    expected_keyword_part = ATS_WEIGHTS["keywords"] * NEUTRAL_KEYWORD_SCORE
    assert report.score == pytest.approx(
        expected_keyword_part
        + ATS_WEIGHTS["sections"] * 0.0
        + ATS_WEIGHTS["title_alignment"] * NEUTRAL_TITLE_ALIGNMENT
        + ATS_WEIGHTS["format_safety"] * NEUTRAL_FORMAT_SAFETY,
        abs=0.05,
    )
    assert any("could not be measured" in note for note in report.notes)
    # Sections missing are reported, never invented.
    assert report.section_completeness["education"] is False
    assert set(report.section_completeness) == set(REQUIRED_SECTIONS)


def test_title_alignment_uses_aliases() -> None:
    job = JobListing(
        id="j",
        source="greenhouse",
        title="AI Engineer",
        company="Acme",
    )
    profile = CandidateProfile(
        personal_information=PersonalInformation(
            name="Ahmad Hassan", title="Artificial Intelligence Engineer"
        ),
        skills=["python"],
    )
    report = analyze_ats(TAILORED, profile, job)
    assert report.title_alignment == 100.0


def test_title_alignment_penalises_a_seniority_gap() -> None:
    job = JobListing(id="j", source="greenhouse", title="Senior AI Engineer", company="Acme")
    report = analyze_ats(TAILORED, PROFILE, job)
    assert 0.0 < report.title_alignment < 100.0


def test_notes_never_propose_skills_the_profile_lacks() -> None:
    job = JobListing(
        id="j",
        source="greenhouse",
        title="AI Engineer",
        company="Acme",
        # "docker" is in the profile but not surfaced in the tailored resume;
        # "kubernetes" is not in the profile at all.
        required_skills=["kubernetes", "docker"],
        preferred_skills=["terraform"],
    )
    report = analyze_ats(TAILORED, PROFILE, job)
    notes = " ".join(report.notes).lower()

    # No note instructs the candidate to add a skill.
    for verb in ("add", "include", "insert", "mention", "highlight", "put "):
        assert verb not in notes
    # The unsupported keyword is reported as unsupported, and the supported one
    # is reported as merely not surfaced.
    assert "kubernetes — not supported by the candidate profile" in notes
    assert any("not surfaced in the tailored resume" in note for note in report.notes)
    assert "terraform" in notes


def test_analyzer_is_deterministic() -> None:
    first = analyze_ats(TAILORED, PROFILE, JOB)
    second = analyze_ats(TAILORED, PROFILE, JOB)
    assert first.model_dump() == second.model_dump()


def test_format_risks_are_detected_on_a_bad_document(tmp_path: Path) -> None:
    path = _bad_docx(tmp_path / "bad.docx")
    report = analyze_ats(TAILORED, PROFILE, JOB, docx_path=path)
    risks = " ".join(report.format_risks).lower()
    assert "tables used for layout" in risks
    assert "text boxes detected" in risks
    assert "image(s) present" in risks
    assert "missing standard section heading" in risks
    assert "empty section" in risks
    assert "font sizes" in risks
    assert "no machine-readable contact line" in risks


def test_clean_document_has_no_format_risks(tmp_path: Path) -> None:
    path = _clean_docx(tmp_path / "clean.docx")
    report = analyze_ats(TAILORED, PROFILE, JOB, docx_path=path)
    assert report.format_risks == []
    assert report.score == pytest.approx(100.0, abs=0.05)


def test_format_risks_lower_the_score(tmp_path: Path) -> None:
    clean = analyze_ats(TAILORED, PROFILE, JOB, docx_path=_clean_docx(tmp_path / "c.docx"))
    bad = analyze_ats(TAILORED, PROFILE, JOB, docx_path=_bad_docx(tmp_path / "b.docx"))
    assert bad.score < clean.score
    assert bad.score <= 100.0 - ATS_WEIGHTS["format_safety"] * FORMAT_RISK_PENALTY + 0.05


def test_unreadable_file_is_reported_not_raised(tmp_path: Path) -> None:
    broken = tmp_path / "broken.docx"
    broken.write_bytes(b"not a docx at all")
    report = analyze_ats(TAILORED, PROFILE, JOB, docx_path=broken)
    assert any("could not inspect" in risk.lower() for risk in report.format_risks)
