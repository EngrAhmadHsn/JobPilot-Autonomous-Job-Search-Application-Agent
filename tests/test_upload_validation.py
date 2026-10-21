"""WS-06 · Phase G — resume upload validation (SEC-10).

An uploaded resume is an untrusted file. These tests pin the rejection paths
(disallowed type, oversize, empty, a renamed ZIP, a non-UTF-8 text file) and
the property that matters most: a hostile *filename* can never decide where
the file is stored.
"""

from __future__ import annotations

import io
import zipfile
from pathlib import Path

import pytest
from docx import Document
from fastapi import UploadFile

from src.api.validation import (
    ALLOWED_EXTENSIONS,
    MAX_UPLOAD_BYTES,
    UploadValidationError,
    validate_resume_upload,
)
from src.config.settings import get_settings


@pytest.fixture(autouse=True)
def isolated_upload_dir(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    """Keep the stored uploads out of the developer's data directory."""
    target = tmp_path / "uploads"
    monkeypatch.setattr(get_settings(), "resume_upload_dir", str(target))
    return target


def _upload(name: str, payload: bytes) -> UploadFile:
    return UploadFile(filename=name, file=io.BytesIO(payload))


def _docx_bytes() -> bytes:
    document = Document()
    document.add_paragraph("Ada Lovelace")
    document.add_paragraph("PROFESSIONAL SUMMARY")
    document.add_paragraph("Engineer.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class TestAcceptedUploads:
    def test_a_real_docx_is_stored(self, isolated_upload_dir: Path) -> None:
        path = validate_resume_upload(_upload("resume.docx", _docx_bytes()))

        assert path.exists()
        assert path.parent == isolated_upload_dir
        assert path.suffix == ".docx"

    def test_a_utf8_text_resume_is_stored(self) -> None:
        path = validate_resume_upload(_upload("resume.txt", b"Ada Lovelace\nEngineer"))

        assert path.exists()
        assert path.suffix == ".txt"

    def test_a_markdown_resume_is_stored(self) -> None:
        path = validate_resume_upload(_upload("resume.md", b"# Ada Lovelace"))

        assert path.exists()
        assert path.suffix == ".md"

    def test_the_stored_name_is_generated_not_uploaded(self, isolated_upload_dir: Path) -> None:
        path = validate_resume_upload(_upload("../../evil.docx", _docx_bytes()))

        assert path.parent == isolated_upload_dir
        assert "evil" not in path.name
        assert path.name == f"{path.stem}.docx"


class TestRejectedUploads:
    @pytest.mark.parametrize("name", ["resume.pdf", "resume.exe", "resume", "resume.doc"])
    def test_disallowed_extensions_are_rejected(self, name: str) -> None:
        with pytest.raises(UploadValidationError):
            validate_resume_upload(_upload(name, b"payload"))

    def test_an_oversized_file_is_rejected(self) -> None:
        payload = b"x" * (MAX_UPLOAD_BYTES + 1)

        with pytest.raises(UploadValidationError, match="exceeds"):
            validate_resume_upload(_upload("resume.txt", payload))

    def test_an_empty_file_is_rejected(self) -> None:
        with pytest.raises(UploadValidationError, match="empty"):
            validate_resume_upload(_upload("resume.txt", b""))

    def test_a_renamed_zip_is_rejected(self) -> None:
        """A ZIP that is not a document must not pass as a DOCX."""
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            archive.writestr("hello.txt", "not a document")
        payload = buffer.getvalue()

        with pytest.raises(UploadValidationError):
            validate_resume_upload(_upload("resume.docx", payload))

    def test_a_corrupt_docx_is_rejected(self) -> None:
        with pytest.raises(UploadValidationError):
            validate_resume_upload(_upload("resume.docx", b"PK\x03\x04broken"))

    def test_a_non_utf8_text_file_is_rejected(self) -> None:
        with pytest.raises(UploadValidationError, match="UTF-8"):
            validate_resume_upload(_upload("resume.txt", b"\xff\xfe\x00binary"))

    def test_nothing_is_stored_when_the_upload_is_rejected(self, isolated_upload_dir: Path) -> None:
        with pytest.raises(UploadValidationError):
            validate_resume_upload(_upload("resume.pdf", b"payload"))

        stored = list(isolated_upload_dir.glob("*")) if isolated_upload_dir.exists() else []
        assert stored == []


class TestConstants:
    def test_the_allowed_types(self) -> None:
        assert ALLOWED_EXTENSIONS == frozenset({".docx", ".txt", ".md"})

    def test_the_size_cap_is_five_megabytes(self) -> None:
        assert MAX_UPLOAD_BYTES == 5 * 1024 * 1024
