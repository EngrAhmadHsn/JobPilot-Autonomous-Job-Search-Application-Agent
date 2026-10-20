"""Render a tailored resume to a DOCX deliverable (WS-05).

The renderer writes **only** content that exists: fields are copied verbatim
from the candidate profile or from the already validated tailored content.
It never completes, "polishes" or invents a fact.

The document layout is chosen so that the existing parser
(``src/resume/parser.py``) reads the file back: the same section headings,
the same "Title | Company" experience line, the same "Category: a, b" skill
line. That round-trip is the strongest anti-fabrication check available on a
renderer — anything it invented would have to survive a re-parse to stay
invisible, and it is locked by ``tests/test_render.py``.
"""

from __future__ import annotations

import hashlib
import logging
import re
from pathlib import Path

from src.core.text import normalize
from src.resume.profile import CandidateProfile
from src.resume.tailor import TailoredResume

logger = logging.getLogger(__name__)

#: Section headings, spelled exactly as the resume parser recognizes them.
HEADING_SUMMARY = "Professional Summary"
HEADING_SKILLS = "Technical Skills"
HEADING_EXPERIENCE = "Experience"
HEADING_PROJECTS = "Projects"
HEADING_EDUCATION = "Education"
HEADING_CERTIFICATIONS = "Certifications"
HEADING_ACHIEVEMENTS = "Achievements"

#: Characters that are illegal (or awkward) in a file name on the supported
#: platforms. Job ids look like "greenhouse:abc123", so this matters.
_UNSAFE_FILENAME_CHARS = re.compile(r"[^A-Za-z0-9_.-]")


def _safe_component(value: str) -> str:
    """A file-name-safe version of an id."""
    return _UNSAFE_FILENAME_CHARS.sub("_", value)


def tailored_docx_path(profile_id: str, job_id: str, out_dir: Path) -> Path:
    """Deterministic path of the rendered resume for one (profile, job) pair.

    The name is ``{profile_id}_{safe_job_id}.docx``. When sanitising changed
    the job id (its ``:`` is illegal on Windows), a short digest of the raw id
    is appended so two different job ids can never collide on one file. The
    same inputs always yield the same path; re-running the pipeline replaces
    the file.
    """
    safe_job = _safe_component(job_id)
    if safe_job != job_id:
        digest = hashlib.sha256(job_id.encode("utf-8")).hexdigest()[:8]
        safe_job = f"{safe_job}_{digest}"
    safe_profile = _safe_component(profile_id)
    return Path(out_dir) / f"{safe_profile}_{safe_job}.docx"


def _contact_line(profile: CandidateProfile) -> str:
    """The contact line, built from verbatim profile fields only."""
    info = profile.personal_information
    parts = [
        part
        for part in (info.email, info.phone, info.location, info.website)
        if part and part.strip()
    ]
    return " | ".join(part.strip() for part in parts)


def _skill_lines(tailored: TailoredResume, profile: CandidateProfile) -> list[str]:
    """Skills grouped by the profile's own categories, tailored selection only.

    A category is emitted only when at least one of its skills survived the
    tailoring; skills the candidate does not have are never added.
    """
    highlighted = {normalize(skill) for skill in tailored.highlighted_skills}
    if not highlighted:
        return []

    lines: list[str] = []
    for category, items in profile.skill_categories.items():
        selected = [item for item in items if normalize(item) in highlighted]
        if selected:
            lines.append(f"{category}: {', '.join(selected)}")

    if not lines:
        lines.append(f"Skills: {', '.join(tailored.highlighted_skills)}")
    return lines


def _dates_for(profile: CandidateProfile, company: str) -> str:
    """The dates of the profile entry with this company (verbatim, or empty)."""
    for entry in profile.experience:
        if entry.company and normalize(entry.company) == normalize(company):
            return entry.dates
    return ""


def _project_name(name: str) -> str:
    """A project name the parser cannot split.

    The parser treats " — " and tabs inside a project line as name/subtitle
    and name/tech separators, which would truncate the name. A plain hyphen
    keeps the name intact and machine-readable.
    """
    return name.replace("\t", " ").replace(" — ", " - ").replace(" – ", " - ")


def render_tailored_docx(
    tailored: TailoredResume, profile: CandidateProfile, out_path: Path
) -> Path:
    """Render a tailored resume to a DOCX file.

    Args:
        tailored: The validated tailored content.
        profile: The candidate profile the content came from.
        out_path: Target file; parent directories are created as needed.

    Returns:
        The path of the written file.
    """
    from docx import Document

    document = Document()
    info = profile.personal_information

    if info.name.strip():
        document.add_paragraph(info.name.strip(), style="Title")
    if info.title.strip():
        document.add_paragraph(info.title.strip(), style="Subtitle")
    contact = _contact_line(profile)
    if contact:
        document.add_paragraph(contact)

    if tailored.summary.strip():
        document.add_paragraph(HEADING_SUMMARY, style="Heading 1")
        document.add_paragraph(tailored.summary.strip())

    skill_lines = _skill_lines(tailored, profile)
    if skill_lines:
        document.add_paragraph(HEADING_SKILLS, style="Heading 1")
        for line in skill_lines:
            document.add_paragraph(line)

    if tailored.experience:
        document.add_paragraph(HEADING_EXPERIENCE, style="Heading 1")
        for entry in tailored.experience:
            parts = [entry.title.strip(), entry.company.strip()]
            dates = _dates_for(profile, entry.company)
            if dates:
                parts.append(dates)
            document.add_paragraph(" | ".join(part for part in parts if part))
            for bullet in entry.bullets:
                if bullet.strip():
                    document.add_paragraph(bullet.strip(), style="List Bullet")

    if tailored.projects:
        document.add_paragraph(HEADING_PROJECTS, style="Heading 1")
        for project in tailored.projects:
            document.add_paragraph(_project_name(project.name))
            for bullet in project.bullets:
                if bullet.strip():
                    document.add_paragraph(bullet.strip(), style="List Bullet")

    if profile.education:
        document.add_paragraph(HEADING_EDUCATION, style="Heading 1")
        for entry in profile.education:
            if entry.degree.strip():
                document.add_paragraph(entry.degree.strip())
            if entry.institution.strip():
                document.add_paragraph(entry.institution.strip())
            if entry.dates.strip():
                document.add_paragraph(entry.dates.strip())
            for detail in entry.details:
                if detail.strip():
                    document.add_paragraph(detail.strip(), style="List Bullet")

    if profile.certifications:
        document.add_paragraph(HEADING_CERTIFICATIONS, style="Heading 1")
        for item in profile.certifications:
            if item.strip():
                document.add_paragraph(item.strip(), style="List Bullet")

    if profile.achievements:
        document.add_paragraph(HEADING_ACHIEVEMENTS, style="Heading 1")
        for item in profile.achievements:
            if item.strip():
                document.add_paragraph(item.strip(), style="List Bullet")

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(out_path))
    logger.info("Tailored resume rendered", extra={"path": out_path.name})
    return out_path
